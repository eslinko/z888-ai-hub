"""
Tests for DOC/DOCX document processor.
"""

import os
import pytest
from docx import Document
from z888_ai_hub.processors import DocProcessor
from z888_ai_hub.processors.exceptions import ContentExtractionError

# Константы
SAMPLE_DOCS_DIR = "tests/sample_docs"
OUTPUT_JSON_DIR = "tests/sample_docs/json"

# Фикстуры
@pytest.fixture
def doc_processor():
    return DocProcessor()

@pytest.fixture
def setup_test_files(tmp_path):
    """Create test DOC files."""
    # Создаем простой DOC файл
    doc = Document()
    doc.add_heading('Test Document', 0)
    doc.add_paragraph('This is a test paragraph.')
    doc.add_paragraph('This is another paragraph with some formatting.').runs[0].bold = True
    
    # Добавляем таблицу
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Data 1"
    table.cell(1, 1).text = "Data 2"
    
    # Сохраняем файл
    test_doc_path = os.path.join(tmp_path, "test.docx")
    doc.save(test_doc_path)
    
    return tmp_path

@pytest.mark.asyncio
async def test_doc_content_extraction(doc_processor, setup_test_files):
    """Test basic content extraction from DOC file."""
    test_doc_path = os.path.join(setup_test_files, "test.docx")
    
    # Извлекаем содержимое
    content = await doc_processor.extract_content(test_doc_path)
    
    # Проверяем базовую структуру
    assert content.text, "Text should not be empty"
    assert "Test Document" in content.text, "Heading should be in text"
    assert "test paragraph" in content.text, "Paragraph content should be in text"
    
    # Проверяем метаданные
    assert content.metadata, "Metadata should not be empty"
    assert content.metadata["paragraph_count"] > 0, "Should have paragraphs"
    assert content.metadata["section_count"] > 0, "Should have sections"
    
    # Проверяем таблицы
    assert len(content.tables) == 1, "Should have one table"
    assert content.tables[0]["row_count"] == 2, "Table should have 2 rows"
    assert content.tables[0]["column_count"] == 2, "Table should have 2 columns"
    
    # Проверяем стили
    assert content.styles, "Should have some styles"

@pytest.mark.asyncio
async def test_doc_processing_invalid_file(doc_processor):
    """Test processing of invalid DOC file."""
    with pytest.raises(ContentExtractionError):
        await doc_processor.extract_content("nonexistent.doc")

@pytest.mark.asyncio
async def test_doc_full_processing(doc_processor, setup_test_files):
    """Test full document processing including JSON generation."""
    test_doc_path = os.path.join(setup_test_files, "test.docx")
    
    # Создаем директорию для выходных файлов
    os.makedirs(OUTPUT_JSON_DIR, exist_ok=True)
    
    # Обрабатываем документ
    result = await doc_processor.process_document(test_doc_path, OUTPUT_JSON_DIR)
    
    # Проверяем результат
    assert result.file_id, "Should have file_id"
    assert result.file_name == "test.docx", "Should have correct file name"
    assert result.summary, "Should have generated summary"
    
    # Проверяем созданный JSON файл
    json_path = os.path.join(OUTPUT_JSON_DIR, "test_docx.json")
    assert os.path.exists(json_path), "JSON file should be created"

@pytest.mark.asyncio
async def test_doc_metadata_extraction(doc_processor, setup_test_files):
    """Test metadata extraction from DOC file."""
    test_doc_path = os.path.join(setup_test_files, "test.docx")
    
    # Извлекаем содержимое
    content = await doc_processor.extract_content(test_doc_path)
    
    # Проверяем все поля метаданных
    assert isinstance(content.metadata, dict), "Metadata should be a dictionary"
    assert "title" in content.metadata, "Should have title"
    assert "author" in content.metadata, "Should have author"
    assert "created" in content.metadata, "Should have creation date"
    assert "modified" in content.metadata, "Should have modification date"
    assert "paragraph_count" in content.metadata, "Should have paragraph count"
    assert "section_count" in content.metadata, "Should have section count"
    assert "page_count" in content.metadata, "Should have page count"

@pytest.mark.asyncio
async def test_doc_table_extraction(doc_processor, setup_test_files):
    """Test table extraction from DOC file."""
    test_doc_path = os.path.join(setup_test_files, "test.docx")
    
    # Извлекаем содержимое
    content = await doc_processor.extract_content(test_doc_path)
    
    # Проверяем структуру таблицы
    assert len(content.tables) == 1, "Should have one table"
    table = content.tables[0]
    
    assert table["row_count"] == 2, "Table should have 2 rows"
    assert table["column_count"] == 2, "Table should have 2 columns"
    
    # Проверяем данные таблицы
    assert "Header 1" in table["data"][0][0]["text"], "Should have correct header"
    assert "Data 1" in table["data"][1][0]["text"], "Should have correct data"

@pytest.mark.asyncio
async def test_doc_style_extraction(doc_processor, setup_test_files):
    """Test style extraction from DOC file."""
    test_doc_path = os.path.join(setup_test_files, "test.docx")
    
    # Извлекаем содержимое
    content = await doc_processor.extract_content(test_doc_path)
    
    # Проверяем стили
    assert content.styles, "Should have styles"
    for style_name, style_info in content.styles.items():
        assert isinstance(style_info, dict), "Style info should be a dictionary"
        assert "font" in style_info, "Style should have font information"
        assert "size" in style_info, "Style should have size information"
        assert "bold" in style_info, "Style should have bold information"
        assert "italic" in style_info, "Style should have italic information" 